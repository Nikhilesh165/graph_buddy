from pathlib import Path

import pytest
from docx import Document
from fpdf import FPDF

from app.services.parsing import UnsupportedFileType, parse_file


def test_parse_txt(tmp_path: Path) -> None:
    path = tmp_path / "note.txt"
    path.write_text("hello from a text file", encoding="utf-8")

    result = parse_file(path, "text/plain")

    assert result.text == "hello from a text file"
    assert result.char_count == len(result.text)
    assert result.row_count is None


def test_parse_md(tmp_path: Path) -> None:
    path = tmp_path / "note.md"
    path.write_text("# Heading\n\nSome body text.", encoding="utf-8")

    result = parse_file(path, "text/markdown")

    assert "Heading" in result.text
    assert "Some body text." in result.text


def test_parse_csv(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_text("name,age\nAlice,30\nBob,40\n", encoding="utf-8")

    result = parse_file(path, "text/csv")

    assert result.row_count == 2
    assert "name" in result.text and "Alice" in result.text


def test_parse_docx(tmp_path: Path) -> None:
    path = tmp_path / "doc.docx"
    document = Document()
    document.add_paragraph("Hello from a docx file")
    document.save(path)

    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    result = parse_file(path, docx_mime)

    assert "Hello from a docx file" in result.text


def test_parse_pdf(tmp_path: Path) -> None:
    path = tmp_path / "doc.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Hello from a test PDF")
    pdf.output(str(path))

    result = parse_file(path, "application/pdf")

    assert "Hello from a test PDF" in result.text


def test_parse_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "data.xlsx"
    path.write_bytes(b"not a real xlsx")

    with pytest.raises(UnsupportedFileType):
        parse_file(path, "application/octet-stream")
